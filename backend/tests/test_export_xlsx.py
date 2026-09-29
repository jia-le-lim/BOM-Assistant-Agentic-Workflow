"""The workbook export.

What matters is that a number lands in the cell an engineer would otherwise
have typed it into, and that a row nobody approved stays empty. Both are
asserted by reading the produced file back, not by trusting the writer.
"""

import io

from conftest import ADMIN, ENG, VIEWER, upload
from app.db import get_conn
from openpyxl import load_workbook


def scored_batch(client, synth_csv) -> int:
    b = upload(client, synth_csv).json()["batch_id"]
    r = client.post(f"/run-recommendation?batch_id={b}", headers=ENG)
    assert r.status_code == 200, r.text
    return b


def sheets(payload: bytes):
    wb = load_workbook(io.BytesIO(payload))
    bom = wb["BOM"]
    header = [c.value for c in next(bom.iter_rows(min_row=1, max_row=1))]
    rows = {}
    for r in bom.iter_rows(min_row=2, values_only=True):
        cells = dict(zip(header, r))
        rows[(str(cells["item_id"]), str(cells["stockroom_id"]))] = cells

    st = wb["Review status"]
    st_header = [c.value for c in next(st.iter_rows(min_row=1, max_row=1))]
    status = {}
    for r in st.iter_rows(min_row=2, values_only=True):
        cells = dict(zip(st_header, r))
        status[(str(cells["item_id"]), str(cells["stockroom_id"]))] = cells
    return header, rows, status


def get_xlsx(client, batch_id, headers=ENG):
    r = client.get(f"/export/wings.xlsx?batch_id={batch_id}", headers=headers)
    assert r.status_code == 200, r.text
    return r


def seed_historical_approval(batch_id):
    """Existing approvals remain exportable after workspaces become private."""
    conn = get_conn()
    try:
        conn.execute("UPDATE review_history SET senior_approved_by=?, "
                     "senior_approved_at=datetime('now') WHERE batch_id=?",
                     ("historical-approver", batch_id))
        conn.commit()
    finally:
        conn.close()


def test_approved_values_land_in_their_columns(client, synth_csv, db_file):
    b = scored_batch(client, synth_csv)
    # 100005 is high risk; simulate an approval recorded before isolation.
    client.post(f"/review/100005?batch_id={b}",
                json={"decision": "override", "final_max": 7, "final_rop": 4,
                      "final_min": 2, "justification": "Constraint tool"},
                headers=ENG)
    seed_historical_approval(b)

    r = get_xlsx(client, b)
    _header, rows, status = sheets(r.content)

    row = rows[("100005", "24")]
    assert row["factory_recommended_new_max"] == 7
    assert row["factory_recommended_new_rop"] == 4
    assert row["factory_recommended_new_min"] == 2
    assert row["review_acknowledge"] == "Y"
    assert row["modified_user"] == "alice"
    assert row["modified_date"]                       # reviewed_at stamp
    assert row["justification"] == "Constraint tool"
    assert status[("100005", "24")]["status"] == "reviewed"
    assert r.headers["X-Rows-Updated"] == "1"


def test_unreviewed_rows_are_left_empty(client, synth_csv, db_file):
    b = scored_batch(client, synth_csv)
    r = get_xlsx(client, b)
    _header, rows, status = sheets(r.content)

    row = rows[("100005", "24")]
    assert row["factory_recommended_new_max"] in (None, "")
    assert row["review_acknowledge"] in (None, "")
    assert status[("100005", "24")]["status"] == "pending_review"
    assert r.headers["X-Rows-Updated"] == "0"
    assert int(r.headers["X-Pending-Review"]) > 0


def test_awaiting_senior_is_not_exported(client, synth_csv, db_file):
    """A high-risk accept is a decision, but not an approved one."""
    b = scored_batch(client, synth_csv)
    res = client.post(f"/review/100005?batch_id={b}",
                      json={"decision": "accept"}, headers=ENG).json()
    assert res["requires_senior_approval"] is True

    r = get_xlsx(client, b)
    _header, rows, status = sheets(r.content)
    assert rows[("100005", "24")]["factory_recommended_new_max"] in (None, "")
    assert status[("100005", "24")]["status"] == "awaiting_senior"
    assert r.headers["X-Awaiting-Senior"] == "1"


def test_a_review_that_changes_nothing_is_acknowledged_not_updated(
        client, synth_csv, db_file):
    """Writing the same numbers back would present a no-op as a stock update."""
    b = scored_batch(client, synth_csv)
    # 100008: engine proposes 2 -> 3; rejecting keeps the current values.
    client.post(f"/review/100008?batch_id={b}",
                json={"decision": "reject"}, headers=ENG)

    r = get_xlsx(client, b)
    _header, rows, status = sheets(r.content)
    row = rows[("100008", "24")]
    assert row["review_acknowledge"] == "Y"
    assert row["factory_recommended_new_max"] in (None, "")
    assert status[("100008", "24")]["note"] == "reviewed, no change"
    assert r.headers["X-Rows-Acknowledged"] == "1"
    assert r.headers["X-Rows-Updated"] == "0"


def test_the_sheet_keeps_the_input_columns_and_values(client, synth_csv, db_file):
    b = scored_batch(client, synth_csv)
    header, rows, _status = sheets(get_xlsx(client, b).content)

    uploaded = synth_csv.decode().splitlines()[0].split(",")
    assert header[:len(uploaded)] == uploaded          # same columns, same order
    assert rows[("100005", "24")]["unitprice"] == "5000"
    assert rows[("100005", "24")]["item_desc"] == "TEST PART"


def test_quarantined_rows_are_present_but_flagged(client, synth_csv, db_file):
    b = scored_batch(client, synth_csv)
    _header, rows, status = sheets(get_xlsx(client, b).content)
    # 100003 has negative consumption; ingestion quarantines it.
    assert ("100003", "24") in rows
    assert status[("100003", "24")]["status"] == "quarantined"
    assert status[("100003", "24")]["note"] == "NEGATIVE_CONSUMPTION"


def test_input_commentary_survives_on_rows_nobody_reviewed(
        client, synth_csv, db_file):
    """Sheet 1 promises the input back as it was read.

    justification/comments are commentary, not approval signals, so clearing
    them on every row discarded whatever WINGS sent on the pending, cleared and
    quarantined rows -- which is most of the file.
    """
    b = scored_batch(client, synth_csv)
    _header, rows, status = sheets(get_xlsx(client, b).content)

    # make_row() seeds both columns; an unreviewed row must still carry them.
    assert status[("100005", "24")]["status"] == "pending_review"
    assert rows[("100005", "24")]["justification"] == "LEAK_IF_USED"
    assert rows[("100005", "24")]["comments"] == "LEAK_IF_USED"

    # The approval signals are still cleared unconditionally -- that is the
    # stale-approval control and it must not have been relaxed.
    assert rows[("100005", "24")]["factory_recommended_new_max"] in (None, "")
    assert rows[("100005", "24")]["review_acknowledge"] in (None, "")
    assert rows[("100005", "24")]["modified_user"] in (None, "")

    # A reviewed row's text is replaced by the review's own.
    client.post(f"/review/100008?batch_id={b}",
                json={"decision": "reject", "justification": "keep as is"},
                headers=ENG)
    _header, rows, _status = sheets(get_xlsx(client, b).content)
    assert rows[("100008", "24")]["justification"] == "keep as is"


def test_reviewer_text_cannot_become_a_live_formula(client, synth_csv, db_file):
    """The workbook is emailed and re-imported, so a justification starting with
    '=' would execute when the next person opens it."""
    b = scored_batch(client, synth_csv)
    client.post(f"/review/100005?batch_id={b}",
                json={"decision": "override", "final_max": 7, "final_rop": 4,
                      "final_min": 2,
                      "justification": '=HYPERLINK("http://x","click")',
                      "comment": "-1+1"},
                headers=ENG)
    seed_historical_approval(b)

    _header, rows, _status = sheets(get_xlsx(client, b).content)
    row = rows[("100005", "24")]
    assert row["justification"].startswith("'="), row["justification"]
    assert row["comments"].startswith("'-"), row["comments"]
    # The approved numbers are ints, not text, so they are untouched.
    assert row["factory_recommended_new_max"] == 7


def test_export_is_role_gated(client, synth_csv, db_file):
    b = scored_batch(client, synth_csv)
    assert client.get(f"/export/wings.xlsx?batch_id={b}",
                      headers=VIEWER).status_code == 403
    assert client.get(f"/export/wings.xlsx?batch_id={b}",
                      headers={**ADMIN, "X-User": "alice"}).status_code == 200


def test_unknown_and_unscored_batches_are_refused(client, synth_csv, db_file):
    assert client.get("/export/wings.xlsx?batch_id=9999",
                      headers=ENG).status_code == 404
    b = upload(client, synth_csv).json()["batch_id"]      # loaded, never scored
    assert client.get(f"/export/wings.xlsx?batch_id={b}",
                      headers=ENG).status_code == 409
